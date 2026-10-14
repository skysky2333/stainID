"""Tables from the files labs keep study information in: Excel workbooks (.xlsx), Word documents (.docx tables), CSV or
tab-separated files and text pasted from a spreadsheet. Every table becomes rows of cell texts; nothing is assumed
about headers, so the map and donor importers decide what each row is."""
from __future__ import annotations

import csv
import datetime
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from xml.etree import ElementTree

PICTURES = {".png", ".jpg", ".jpeg", ".heic", ".bmp", ".gif", ".tif", ".tiff", ".pdf", ".webp"}
WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


@dataclass
class Table:
    name: str          # what to call it: the sheet name, the heading above a Word table, or the file name
    rows: list[list[str]]
    source: str = ""   # the file it came from
    notes: list[str] = field(default_factory=list)


def read_tables(filename: str, content: bytes) -> list[Table]:
    """Every table in a file. Raises ValueError, in plain words, for files that cannot be read (pictures, old formats)."""
    suffix = Path(filename).suffix.lower()
    stem = Path(filename).stem
    if suffix in (".xlsx", ".xlsm"):
        tables = _excel(stem, content)
    elif suffix == ".docx":
        tables = _word(stem, content)
    elif suffix in (".csv", ".tsv", ".txt", ""):
        tables = [Table(stem, parse_text(decode(content)))]
    elif suffix == ".xls":
        raise ValueError(f"{filename} is an old Excel file (.xls): open it in Excel and save it as .xlsx (or CSV), then add that file")
    elif suffix == ".doc":
        raise ValueError(f"{filename} is an old Word file (.doc): open it in Word and save it as .docx, then add that file")
    elif suffix in PICTURES:
        raise ValueError(f"{filename} is a picture, which stainID cannot read: type the map into the grid instead "
                         "(or copy it into a spreadsheet and paste it)")
    else:
        raise ValueError(f"{filename}: stainID reads Excel (.xlsx), Word (.docx) and CSV files, or text pasted from a spreadsheet")
    for table in tables:
        table.source = filename
    found = [t for t in tables if t.rows]
    if not found:
        raise ValueError(f"No table was found in {filename}")
    return found


def decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("latin-1")


def parse_text(text: str) -> list[list[str]]:
    """Rows of a CSV file or of cells pasted from a spreadsheet (tab-separated; quoted cells may hold line breaks)."""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    if not text.strip():
        return []
    if "\t" in text:
        delimiter = "\t"
    else:
        try:
            delimiter = csv.Sniffer().sniff(text[:4096], delimiters=",;|").delimiter
        except csv.Error:
            delimiter = ","
    return tidy([list(row) for row in csv.reader(io.StringIO(text), delimiter=delimiter)])


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value.strftime("%Y-%m-%d")
    return str(value).strip()


def tidy(rows: list[list[object]]) -> list[list[str]]:
    """Cell texts with the empty rows and columns around the table removed, every row as wide as the widest."""
    rows = [[cell_text(c) for c in row] for row in rows]
    while rows and not any(rows[-1]):
        rows.pop()
    while rows and not any(rows[0]):
        rows.pop(0)
    if not rows:
        return []
    width = max((max((i + 1 for i, c in enumerate(r) if c), default=0) for r in rows), default=0)
    rows = [(r + [""] * width)[:width] for r in rows]
    first = min((next(i for i, c in enumerate(r) if c) for r in rows if any(r)), default=0)
    return [r[first:] for r in rows]


def _excel(stem: str, content: bytes) -> list[Table]:
    try:
        from openpyxl import load_workbook
    except ImportError as error:  # pragma: no cover - openpyxl comes with the app extra
        raise ValueError("Reading Excel files needs the openpyxl package (pip install openpyxl); or save the sheet as CSV") from error
    try:
        book = load_workbook(io.BytesIO(content), data_only=True)
    except Exception as error:
        raise ValueError(f"This Excel file could not be opened ({error})") from error
    sheets = [s for s in book.worksheets if s.sheet_state == "visible"] or book.worksheets
    return [Table(sheet.title if len(sheets) > 1 else stem, tidy([list(r) for r in sheet.iter_rows(values_only=True)])) for sheet in sheets]


def _paragraph(paragraph: ElementTree.Element) -> str:
    parts = []
    for element in paragraph.iter():
        if element.tag == WORD + "t":
            parts.append(element.text or "")
        elif element.tag == WORD + "tab":
            parts.append(" ")
        elif element.tag in (WORD + "br", WORD + "cr"):
            parts.append("\n")
    return "".join(parts)


def _word_rows(table: ElementTree.Element) -> list[list[str]]:
    rows = []
    for row in table.findall(WORD + "tr"):
        cells = []
        for cell in row.findall(WORD + "tc"):
            cells.append("\n".join(line for line in (_paragraph(p).strip() for p in cell.findall(WORD + "p")) if line))
            span = cell.find(f"{WORD}tcPr/{WORD}gridSpan")
            cells += [""] * (int(span.get(WORD + "val", "1")) - 1 if span is not None else 0)
        rows.append(cells)
    return tidy(rows)


def _word(stem: str, content: bytes) -> list[Table]:
    try:
        document = zipfile.ZipFile(io.BytesIO(content)).read("word/document.xml")
    except (zipfile.BadZipFile, KeyError) as error:
        raise ValueError(f"This Word file could not be opened ({error})") from error
    body = ElementTree.fromstring(document).find(WORD + "body")
    tables, heading = [], ""
    for element in body if body is not None else []:
        if element.tag == WORD + "p":
            text = re.sub(r"\s+", " ", _paragraph(element)).strip()
            heading = text or heading
        elif element.tag == WORD + "tbl":
            tables.append(Table(heading or stem, _word_rows(element)))
            heading = ""
    if len(tables) > 1 and all(t.name == stem for t in tables):
        for number, table in enumerate(tables, start=1):
            table.name = f"{stem} (table {number})"
    return tables
